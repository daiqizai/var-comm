"""Native Word tables from the frozen table CSVs, independent of Pandoc tabular parsing."""
from pathlib import Path
import csv
import hashlib
import re

from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def braced_argument(text, start):
    while start < len(text) and text[start].isspace():
        start += 1
    if start >= len(text) or text[start] != '{':
        raise ValueError('Expected a braced TeX argument')
    depth = 1
    end = start + 1
    while depth and end < len(text):
        if text[end] == '\\':
            end += 2
            continue
        if text[end] == '{':
            depth += 1
        elif text[end] == '}':
            depth -= 1
        end += 1
    if depth:
        raise ValueError('Unclosed TeX argument')
    return text[start + 1:end - 1], end


def plain_caption(text):
    replacements = {
        r'\%': '%', r'\uparrow': '↑', r'\downarrow': '↓',
        r'\Delta': 'Δ', r'\to': '→', r'\footnotesize': '',
        r'\smallskip': '', r'\par': ' ', r'\,': ' ', r'\;': ' ',
        '~': ' ', '$': '',
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r'\s+', ' ', text).strip()
    if '\\' in text or '{' in text or '}' in text:
        raise ValueError('Unconverted TeX in table caption or note: ' + text)
    return text


def table_spec(tex, number):
    marker = 'NESTNATIVETABLE' + ('ONE', 'TWO', 'THREE')[number - 1]
    cap_start = tex.index(r'\caption') + len(r'\caption')
    caption, _ = braced_argument(tex, cap_start)
    note = ''
    if r'\parbox' in tex:
        start = tex.index(r'\parbox') + len(r'\parbox')
        _, end = braced_argument(tex, start)
        note, _ = braced_argument(tex, end)
    elif r'\begin{minipage}' in tex:
        start = tex.index(r'\begin{minipage}') + len(r'\begin{minipage}')
        _, end = braced_argument(tex, start)
        note = tex[end:tex.index(r'\end{minipage}', end)]
    return dict(number=number, marker=marker,
                caption=plain_caption(caption), note=plain_caption(note))


def read_csv(path):
    with path.open(newline='', encoding='utf8') as stream:
        return list(csv.DictReader(stream))


def unique(rows, **keys):
    match = [row for row in rows if all(row[key] == str(value) for key, value in keys.items())]
    if len(match) != 1:
        raise ValueError((keys, len(match)))
    return match[0]


def paired_cell(row, multiline=False):
    mean = f"{float(row['mean']):+.4f}"
    bounds = f"[{float(row['ci_low']):.4f}, {float(row['ci_high']):.4f}]"
    return mean + ('\n' if multiline else ' ') + bounds


def data_rows(number, rows):
    """Yield (kind, texts), preserving the TeX table block and data order."""
    if number == 1:
        assert len(rows) == 11
        yield 'group', ['(A) Main policies at N = 1024']
        yield 'header', ['SNR', 'Util. W→P (%)', 'ΔLPIPS [95% CI]']
        for snr in (4, 10, 19):
            row = unique(rows, block='A', population='common500', snr_db=snr)
            utilization = (f"{float(row['utilization_W_percent']):.2f} → "
                           f"{float(row['utilization_P_percent']):.2f}")
            yield 'data', [str(snr), utilization, paired_cell(row)]
        yield 'group', ['(B) Fixed-granularity diagnostic, 10 dB']
        yield 'header', ['Grid (candidates)', 'Utilization (%)', 'LPIPS']
        for arm in ('WHOLE', 'HALF', 'QUARTER', 'EIGHTH'):
            row = unique(rows, block='B', arm=arm, snr_db=10)
            yield 'data', [f"{arm.title()} ({int(row['candidates'])})",
                           f"{float(row['utilization_percent']):.2f}",
                           f"{float(row['mean']):.4f}"]
        yield 'group', ['(C) Receiver interventions: ΔLPIPS [95% CI]']
        yield 'header', ['SNR', 'B0−A0', 'C0−B0']
        for snr in (10, 19):
            left = unique(rows, block='C', snr_db=snr, comparison='B0_minus_A0')
            right = unique(rows, block='C', snr_db=snr, comparison='C0_minus_B0')
            yield 'data', [str(snr), paired_cell(left, True), paired_cell(right, True)]
    elif number == 2:
        assert len(rows) == 18
        yield 'header', ['Method', '4 dB', '10 dB', '19 dB']
        for method in dict.fromkeys(row['method'] for row in rows):
            cells = []
            for snr in (4, 10, 19):
                row = unique(rows, method=method, snr_db=snr)
                cells.append(f"{float(row['psnr_db']):.3f} / {float(row['clip_image_cosine']):.3f}")
            yield 'data', [row['label'], *cells]
    elif number == 3:
        assert len(rows) == 12
        yield 'header', ['SNR', 'ΔLPIPS', 'ΔDreamSim']
        for block, comparison, label in (
            ('A', 'NEST_minus_WHOLE', 'NeST-Com minus WHOLE'),
            ('B', 'NEST_minus_ADA', 'NeST-Com minus Ada (adapted)'),
        ):
            yield 'group', [f'({block}) {label}']
            for snr in (4, 10, 19):
                left = unique(rows, comparison=comparison, snr_db=snr, metric='lpips_alex')
                right = unique(rows, comparison=comparison, snr_db=snr, metric='dreamsim')
                yield 'data', [str(snr), paired_cell(left, True), paired_cell(right, True)]
    else:
        raise ValueError(number)


def paragraph_format(paragraph, *, bold=False, left=False, keep_next=True):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.0
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_together = True
    fmt.keep_with_next = keep_next
    fmt.widow_control = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = bold


def fill_cell(cell, text, *, kind, left=False):
    cell.text = text
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for paragraph in cell.paragraphs:
        paragraph_format(paragraph, bold=kind != 'data', left=left)
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for edge, amount in [('top', 40), ('bottom', 40), ('left', 65), ('right', 65)]:
        item = OxmlElement('w:' + edge)
        item.set(qn('w:w'), str(amount))
        item.set(qn('w:type'), 'dxa')
        margins.append(item)
    props.append(margins)
    if kind in ('group', 'header'):
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'EDEDED' if kind == 'group' else 'F5F5F5')
        shading.set(qn('w:val'), 'clear')
        props.append(shading)


def insert_csv_tables(document, manuscript, specs):
    records = []
    widths_by_number = {1: [1.5, 2.5, 2.5], 2: [2.0, 1.5, 1.5, 1.5], 3: [0.7, 2.9, 2.9]}
    assert len(specs) == 3
    for spec in specs:
        number = spec['number']
        source = Path(manuscript) / 'tables' / ('table_' + ('i', 'ii', 'iii')[number - 1] + '.csv')
        values = read_csv(source)
        content = list(data_rows(number, values))
        markers = [p for p in document.paragraphs if p.text.strip() == spec['marker']]
        if len(markers) != 1:
            raise ValueError(('Table placeholder count', spec['marker'], len(markers)))
        anchor = markers[0]._p
        caption = document.add_paragraph(spec['caption'], style='Caption')
        paragraph_format(caption, left=True)
        caption.paragraph_format.space_before = Pt(6)
        caption.paragraph_format.space_after = Pt(4)
        anchor.addprevious(caption._p)

        widths = widths_by_number[number]
        assert sum(widths) == 6.5
        table = document.add_table(rows=0, cols=len(widths))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        props = table._tbl.tblPr
        width = props.find(qn('w:tblW'))
        width.set(qn('w:w'), '9360')
        width.set(qn('w:type'), 'dxa')
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
            item = OxmlElement('w:' + edge)
            item.set(qn('w:val'), 'single')
            item.set(qn('w:sz'), '4')
            item.set(qn('w:color'), 'D9D9D9')
            borders.append(item)
        props.append(borders)
        for kind, texts in content:
            row = table.add_row()
            row_props = row._tr.get_or_add_trPr()
            row_props.append(OxmlElement('w:cantSplit'))
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
            if kind == 'group':
                cell = row.cells[0].merge(row.cells[-1])
                fill_cell(cell, texts[0], kind=kind, left=True)
            else:
                assert len(texts) == len(widths)
                for index, (cell, text) in enumerate(zip(row.cells, texts)):
                    fill_cell(cell, text, kind=kind, left=number == 2 and index == 0)
        anchor.addprevious(table._tbl)
        if spec['note']:
            note = document.add_paragraph(spec['note'])
            paragraph_format(note, left=True, keep_next=False)
            note.paragraph_format.space_before = Pt(4)
            note.paragraph_format.space_after = Pt(8)
            anchor.addprevious(note._p)
        else:
            for cell in table.rows[-1].cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = False
        anchor.getparent().remove(anchor)
        records.append(dict(number=number, source_csv=str(source.relative_to(manuscript)),
                            source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                            source_records=len(values), word_rows=len(table.rows),
                            word_columns=len(table.columns), width_inches=sum(widths),
                            caption=spec['caption'], note=spec['note'],
                            cells=[dict(kind=kind, text=texts) for kind, texts in content]))
    assert len(document.tables) == 3
    return records
