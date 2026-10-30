"""Deterministic Word review export from the same LaTeX manuscript text.

Figures are rendered from the frozen vector/photo PDFs for Word embedding only.
Equations become editable OMML through the official portable Pandoc converter.
"""
from pathlib import Path
import hashlib,json,re,subprocess,sys,os,shutil,tempfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from word_tables import table_spec, insert_csv_tables
P=Path(__file__).resolve().parents[1];W=P.parents[2]
QA=Path(os.environ.get('NEST_DOC_WORK',str(Path(tempfile.gettempdir())/'nestcom_wcl_word_export')));QA.mkdir(parents=True,exist_ok=True)
PANDOC=Path(os.environ.get('PANDOC',shutil.which('pandoc') or str(W/'.research/kodak_current_20261026_v1/document_tools/pandoc.exe')))
POPPLER=Path(os.environ.get('PDFTOPPM',shutil.which('pdftoppm') or 'C:/Users/11946/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe'))
for i in range(1,4):
    p=P/f'figures/figure{i}.pdf';stem=QA/f'figure{i}'
    if True:  # Render from current PDFs on every document-only build.
        subprocess.run([str(POPPLER),'-png','-singlefile','-r','300',str(p),str(stem)],check=True,capture_output=True)

reference=Document();sec=reference.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
for name in ['Normal','Body Text','First Paragraph','Title','Subtitle','Heading 1','Heading 2','Caption']:
    if name not in reference.styles:continue
    style=reference.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(12);style.paragraph_format.line_spacing=2.0
    style.paragraph_format.space_after=Pt(0);style.paragraph_format.space_before=Pt(0)
reference.styles['Title'].font.size=Pt(16);reference.styles['Title'].font.bold=True
reference.styles['Title'].paragraph_format.line_spacing=1.2
for name in ['Heading 1','Heading 2']:
    reference.styles[name].font.bold=True;reference.styles[name].paragraph_format.keep_with_next=True
    reference.styles[name].paragraph_format.space_before=Pt(8)
reference.styles['Caption'].font.size=Pt(10);reference.styles['Caption'].paragraph_format.line_spacing=1.15
reference.save(QA/'reference.docx')

def flatten(path):
    t=path.read_text(encoding='utf8')
    return re.sub(r'\\input\{([^}]+)\}',lambda m:flatten(P/(m[1]+'.tex')),t)
t=flatten(P/'main.tex')
# TeX-only break hints are unnecessary in Word and prevent inline OMML conversion.
t=t.replace(r'\allowbreak','')
# Word follows the one-column reading order of the same shared sources.
t=re.sub(r'\\ifCLASSOPTIONonecolumn(.*?)\\fi',lambda m:m[1].split(r'\else',1)[0],t,flags=re.S)
expected_display_equations=len(re.findall(r'\\begin\{equation\}',t))
counts={'figure':0,'table':0}
def numbered_caption(m):
    kind='figure' if m[1].startswith('figure') else 'table';counts[kind]+=1
    label=('Fig. '+str(counts[kind])+'. ') if kind=='figure' else ('Table '+('I','II','III')[counts[kind]-1]+'. ')
    return m[0].replace(r'\caption{',r'\caption{'+label,1)
t=re.sub(r'\\begin\{(figure\*?|table)\}(.*?)\\end\{\1\}',numbered_caption,t,flags=re.S)
# Pandoc converts the body, figures and equations only. Complex LaTeX table
# environments are replaced in full and rebuilt as native Word tables below.
table_specs=[]
def native_table_placeholder(m):
    spec=table_spec(m[0],len(table_specs)+1);table_specs.append(spec)
    return '\n\n'+spec['marker']+'\n\n'
t=re.sub(r'\\begin\{table\}.*?\\end\{table\}',native_table_placeholder,t,flags=re.S)
assert len(table_specs)==3 and r'\begin{table}' not in t
# Keep purely numeric confidence intervals as ordinary editable text, copying
# all three number strings without rounding and retaining square brackets.
number_pattern=r'[-+]?\d+(?:\.\d+)?'
numeric_ci_pattern=(r'\$(' + number_pattern + r')\s*\\[,;]\s*\[('
                    + number_pattern + r')\s*,\s*(' + number_pattern + r')\]\$')
t,numeric_ci_text_count=re.subn(numeric_ci_pattern,
    lambda m:(m[1]+' ['+m[2]+', '+m[3]+']').replace('-', '−'),t)
sections=[0,0]
def numbered_heading(m):
    if m[1]=='section':
        sections[0]+=1;sections[1]=0;prefix=('I','II','III','IV')[sections[0]-1]+'. '
    else:
        prefix=chr(65+sections[1])+'. ';sections[1]+=1
    return '\\'+m[1]+'{'+prefix+m[2]+'}'
t=re.sub(r'\\(section|subsection)\{([^}]+)\}',numbered_heading,t)
keys=re.findall(r'\\bibitem\{([^}]+)\}',t);numbers={k:str(i+1)for i,k in enumerate(keys)}
t=re.sub(r'\\cite\{([^}]+)\}',lambda m:'['+', '.join(numbers[k]for k in m[1].split(','))+']',t)
labels={'fig:system':'1','fig:performance':'2','fig:qualitative':'3','tab:internal':'I','tab:external':'II','tab:reduced':'III'}
eqn=[0]
def equation(m):
    eqn[0]+=1;body=m[1]
    for label in re.findall(r'\\label\{([^}]+)\}',body):labels[label]=str(eqn[0])
    body=re.sub(r'\\label\{[^}]+\}','',body)
    return r'\['+body+r'\qquad\text{('+str(eqn[0])+r')}\]'
t=re.sub(r'\\begin\{equation\}(.*?)\\end\{equation\}',equation,t,flags=re.S)
t=re.sub(r'\\eqref\{([^}]+)\}',lambda m:'('+labels[m[1]]+')',t)
t=re.sub(r'\\ref\{([^}]+)\}',lambda m:labels[m[1]],t)
t=re.sub(r'\\IEEEPARstart\{([^}]+)\}\{([^}]+)\}',lambda m:m[1]+m[2],t)
t=t.replace(r'\begin{IEEEkeywords}',r'\paragraph{Index terms}').replace(r'\end{IEEEkeywords}','')
t=t.replace(r'\begin{figure*}[t]',r'\begin{figure}').replace(r'\end{figure*}',r'\end{figure}')
t=t.replace(r'width=\textwidth','width=6.5in')
for i in range(1,4):t=t.replace('{figure'+str(i)+'.pdf}','{'+(QA/f'figure{i}.png').as_posix()+'}')
t=re.sub(r'\\CI\{([^}]+)\}\{([^}]+)\}\{([^}]+)\}',lambda m:'$'+m[1]+'$ [$'+m[2]+', '+m[3]+'$]',t)
t=re.sub(r'\\begin\{minipage\}\{[^}]+\}',r'\n',t).replace(r'\end{minipage}','')
t=t.replace(r'\begingroup','').replace(r'\endgroup','').replace(r'\footnotesize','')
t=re.sub(r'\\begin\{thebibliography\}\{[^}]+\}',lambda m:r'\section*{References}',t)
t=re.sub(r'\\bibitem\{([^}]+)\}',lambda m:'\n\n['+numbers[m[1]]+'] ',t).replace(r'\end{thebibliography}','')
# Text and editable equations come from TeX; native tables use its CSV sources.
source=QA/'review_source.tex';source.write_text(t,encoding='utf8',newline='\n')
target=P/'manuscript_review.docx'
command=[str(PANDOC),str(source),'-f','latex','-t','docx','--standalone','--reference-doc='+str(QA/'reference.docx'),'-o',str(target)]
cp=subprocess.run(command,cwd=P,capture_output=True,text=True,encoding='utf8');(QA/'pandoc.stderr').write_text(cp.stderr,newline='\n');cp.check_returncode()
d=Document(target)
# Pandoc places endChr before sepChr; LibreOffice then resets the right
# delimiter to ')'. Put delimiter properties in their canonical XML order so
# square brackets, norm bars and the empty right side of cases remain exact.
delimiter_order={qn('m:'+name):index for index,name in
                 enumerate(('begChr','sepChr','endChr','grow','shp','ctrlPr'))}
delimiter_properties_reordered=0
for props in d.element.xpath('.//m:dPr'):
    children=list(props)
    ordered=sorted(children,key=lambda child:delimiter_order.get(child.tag,100))
    if children!=ordered:
        for child in children:props.remove(child)
        for child in ordered:props.append(child)
        delimiter_properties_reordered+=1
# Keep punctuation attached to its preceding inline math object. LibreOffice
# otherwise permits a line break immediately before a comma following OMML.
# This changes only Word layout: characters and their order remain identical.
inline_punctuation_attached=0
for paragraph in d.paragraphs:
    children=list(paragraph._p)
    for index,node in enumerate(children[:-1]):
        if node.tag!=qn('m:oMath') or children[index+1].tag!=qn('w:r'):
            continue
        following=children[index+1].find(qn('w:t'))
        if following is None or not following.text or following.text[0] not in ',.;:':
            continue
        punctuation=following.text[0]
        following.text=following.text[1:]
        run=OxmlElement('m:r')
        props=OxmlElement('m:rPr')
        style=OxmlElement('m:sty');style.set(qn('m:val'),'p')
        props.append(style);run.append(props)
        text=OxmlElement('m:t');text.text=punctuation;run.append(text)
        node.append(run)
        inline_punctuation_attached+=1
# Word does not always wrap long reference URLs at slashes. Break opportunities
# affect displayed text only; hyperlink targets and citation contents are intact.
url_wrap_nodes=0
for text_node in d.element.xpath('.//w:t'):
    if text_node.text and ('https://' in text_node.text or 'http://' in text_node.text):
        text_node.text=text_node.text.replace('/', '/\u200b').replace('_', '_\u200b')
        url_wrap_nodes+=1
d.core_properties.title='NeST-Com: Fixed-Budget Next-Scale Token Communication'
d.core_properties.author='Anonymous Authors';d.core_properties.last_modified_by='Anonymous Authors'
d.core_properties.comments=''
for style in d.styles:
    if style.type not in (1,2):continue
    style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr()
    fonts=rp.find(qn('w:rFonts'))
    if fonts is not None:
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme'):
            fonts.attrib.pop(qn('w:'+attr),None)
        for attr in ('ascii','hAnsi','cs'):fonts.set(qn('w:'+attr),'Times New Roman')
    for border in list(style.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    if 'Caption' in style.name:
        style.font.size=Pt(10);style.paragraph_format.line_spacing=1.15
        style.paragraph_format.space_after=Pt(6)
        style.paragraph_format.keep_together=True
    if style.name=='Title':style.font.size=Pt(16);style.font.bold=True
for border in list(d.element.xpath('.//w:pBdr')):border.getparent().remove(border)
for sec in d.sections:
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    p=sec.footer.paragraphs[0];p.alignment=1
    r=OxmlElement('w:r');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r.append(fld);p._p.append(r)
assert not d.tables, 'No Pandoc-parsed table may remain in the Word body'
native_table_records=insert_csv_tables(d,P,table_specs)
for p in d.paragraphs:
    if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
    # Begin the results section with its opening figure and caption together.
    if p.style.name.startswith('Heading') and p.text=='III. Experimental Results':
        p.paragraph_format.page_break_before=True
    if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
    if re.match(r'^(Fig\. [123]\.|Table I{1,3}\.)',p.text):
        p.style=d.styles['Caption'];p.paragraph_format.line_spacing=1.15
        p.paragraph_format.keep_together=True
        if p.text.startswith('Table '):p.paragraph_format.keep_with_next=True
        for run in p.runs:run.font.size=Pt(10);run.bold=False
    p.paragraph_format.widow_control=True
d.save(target)
import zipfile
with zipfile.ZipFile(target)as z:xml=z.read('word/document.xml')
record=dict(source_tex_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),word_sha256=hashlib.sha256(target.read_bytes()).hexdigest(),reference_keys=keys,display_equation_count=eqn[0],expected_display_equations=expected_display_equations,editable_omml_count=xml.count(b'<m:oMath>'),numeric_ci_text_count=numeric_ci_text_count,delimiter_properties_reordered=delimiter_properties_reordered,inline_punctuation_attached=inline_punctuation_attached,native_tables=len(d.tables),native_table_sources=native_table_records,figures=3,body_double_line_spacing=True,figures_unchanged=False, unchanged_figures=[1,3], revised_figures=[2],revision_base_commit='640b359d1c13864184d9d6f0f4a49bb37e622882',figure_pdf_hashes={str(i):hashlib.sha256((P/f'figures/figure{i}.pdf').read_bytes()).hexdigest() for i in range(1,4)},render_review_pending=True)
(QA/'conversion_check.json').write_text(json.dumps(record,indent=2)+'\n',newline='\n')
assert record['display_equation_count']==expected_display_equations and record['editable_omml_count']>=expected_display_equations and record['native_tables']==3
print(json.dumps(record))
